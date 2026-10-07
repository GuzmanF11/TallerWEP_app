# -*- coding: utf-8 -*-
"""Genera Word con preguntas pendientes al cliente — Módulo Compras."""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from pathlib import Path

OUT = Path(__file__).parent / "Preguntas pendientes al cliente - Modulo Compras.docx"


def set_normal_style(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x0E, 0x88, 0xC9)
    return h


def para(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    return p


def bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p


def question_block(doc, num, title, body_lines, note=None):
    p = doc.add_paragraph()
    r_num = p.add_run(f"{num}. ")
    r_num.bold = True
    r_num.font.color.rgb = RGBColor(0x0E, 0x88, 0xC9)
    r_title = p.add_run(title)
    r_title.bold = True
    for line in body_lines:
        bp = doc.add_paragraph(style="List Bullet")
        bp.paragraph_format.left_indent = Cm(1)
        bp.add_run(line)
    if note:
        np = doc.add_paragraph()
        np.paragraph_format.left_indent = Cm(0.5)
        nr = np.add_run(f"Nota: {note}")
        nr.italic = True
        nr.font.size = Pt(10)
        nr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    doc.add_paragraph()


def main():
    doc = Document()
    set_normal_style(doc)

    # Portada / encabezado
    title = doc.add_heading("Preguntas pendientes al cliente", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x0E, 0x88, 0xC9)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Módulo Compras — Taller Web")
    sr.bold = True
    sr.font.size = Pt(14)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run("Fecha: 2 de septiembre de 2026\nElaborado por: equipo de desarrollo Taller Web")
    mr.font.size = Pt(10)
    mr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_paragraph()
    para(
        doc,
        "Este documento reúne todas las decisiones que necesitamos del cliente para continuar "
        "con el módulo de Compras. Las preguntas están ordenadas por prioridad. "
        "Las marcadas como CRÍTICAS desbloquean la Fase 4 (encabezado de factura, totales y estados).",
    )

    doc.add_page_break()

    # --- CRÍTICAS ---
    heading(doc, "A. Preguntas CRÍTICAS — Fase 4 (encabezado y cierre de factura)", 1)
    para(doc, "Sin estas respuestas no podemos guardar ni cerrar un documento de compra completo.", italic=True)

    question_block(
        doc, 1, "Descuento en líneas (Control 28 / 30)",
        [
            "¿El descuento se resta del subtotal o es solo informativo?",
            "En su ejemplo (Acetaminofén 10 %, Batería 5 %) el subtotal no refleja el descuento.",
            "Si es informativo: el costo ya viene descontado y solo guardamos el porcentaje.",
            "Si se aplica: ¿se resta del costo unitario, del subtotal de línea, o de ambos?",
            "¿Guardamos también un campo descuento_monto además del porcentaje?",
        ],
        "Hoy el sistema guarda el % pero no lo resta del subtotal.",
    )

    question_block(
        doc, 2, "Fechas del encabezado de compra (Compras1)",
        [
            "¿Qué fechas lleva el documento?",
            "Fecha de emisión de la factura",
            "Fecha de recepción en bodega",
            "Fecha de ingreso al sistema",
            "¿Cuáles son obligatorias para el control fiscal mensual?",
        ],
    )

    question_block(
        doc, 3, "Estados del documento y cuándo afecta cada cosa",
        [
            "¿Qué estados tiene una compra? (Ej.: Borrador → Confirmada → Anulada)",
            "¿En qué momento la compra afecta el inventario (stock contable / físico)?",
            "¿En qué momento actualiza los costos del producto?",
            "¿En qué momento genera o afecta Cuentas por Pagar?",
            "¿En qué momento entra al reporte fiscal de IVA del mes?",
        ],
        "Es la decisión estructural más importante del módulo.",
    )

    question_block(
        doc, 4, "Moneda y tipo de cambio (importaciones)",
        [
            "¿En qué moneda se captura la factura?",
            "¿Dónde va el tipo de cambio de aduana vs. el costo real pagado al proveedor?",
            "¿Un documento puede mezclar monedas o siempre es una sola?",
        ],
    )

    question_block(
        doc, 5, "Lista exacta de totales a guardar en el encabezado",
        [
            "Total gravado / afecto",
            "Total exento",
            "Total no sujeto",
            "Total XCOMPROBAR (pagos por cuenta ajena)",
            "Total IVA",
            "Total descuentos",
            "Retenciones (¿cuáles aplican en compras?)",
            "Total del documento",
            "Total fiscal (lo que va a Hacienda)",
            "Total contable (lo que va a contabilidad interna)",
            "¿Cuáles se muestran en pantalla y cuáles solo se guardan en base de datos?",
        ],
    )

    question_block(
        doc, 6, "Diferencia entre total fiscal y total contable",
        [
            "El cliente indicó que XCOMPROBAR entra al total del documento pero NO al fiscal.",
            "¿Cuáles son las reglas exactas de cálculo?",
            "¿Hay otros tipos de línea que sigan reglas similares?",
        ],
    )

    doc.add_page_break()

    # --- IMPORTANTES Fases 5-7 ---
    heading(doc, "B. Preguntas IMPORTANTES — Fases 5 a 7", 1)

    question_block(
        doc, 7, "Retaceo de importaciones — gastos a prorratear",
        ["Flete, seguro, DAI/aranceles, almacenaje, transporte interno, otros."],
    )

    question_block(
        doc, 8, "Retaceo — criterio de prorrateo",
        ["Por valor FOB, por peso (libras), por cantidad, por volumen, o combinación."],
    )

    question_block(
        doc, 9, "Retaceo — dónde se guardan los gastos",
        [
            "¿Como líneas XCOMPROBAR en el detalle?",
            "¿En tabla aparte (ej. compras_gastos_importacion)?",
            "¿Ambos?",
        ],
    )

    question_block(
        doc, 10, "Nota de crédito",
        [
            "¿Resta del total fiscal y contable?",
            "¿Revierte inventario y costos automáticamente?",
            "¿Se vincula a la factura original?",
        ],
    )

    question_block(
        doc, 11, "Sujeto excluido",
        [
            "¿Mismas reglas que consumidor final?",
            "¿Tratamiento distinto para IVA e inventario?",
        ],
    )

    question_block(
        doc, 12, "Tabla de Aranceles (SAC de Aduana)",
        [
            "El grid 20 pide DAI y Código arancelario pero hoy no hay origen de datos.",
            "¿El cliente entrega el catálogo SAC?",
            "¿Se carga manual por producto o se busca por código?",
        ],
    )

    question_block(
        doc, 13, "Cuentas por Pagar",
        [
            "El control 12 exige registrar crédito/contado en Cuentas por Pagar.",
            "Los saldos del proveedor (saldo actual, saldo vencido) dependen de ese módulo.",
            "¿Cuándo se implementa? ¿Campos mínimos que necesitan ya?",
            "¿Cómo se calculan saldo_actual y saldo_vencido?",
        ],
    )

    question_block(
        doc, 14, "Bancos y cajas chicas (Control 4)",
        [
            "¿Qué campos necesita cada registro de cuenta bancaria o caja chica?",
            "¿Una compra confirmada siempre afecta una cuenta destino?",
        ],
        "El cliente indicó que la tabla puede crearse ahora y el módulo de control después.",
    )

    doc.add_page_break()

    # --- COSTOS ---
    heading(doc, "C. Costos e inventario — historial y decimales", 1)
    para(doc, "Afectan cuándo la compra actualiza precios y el módulo de Gestión de Costos.", italic=True)

    question_block(
        doc, 15, "Respaldos de costos — contradicción Excel filas 147/149 (BLOQUEANTE)",
        [
            "¿costo_previo1/2 respalda costo_usd?",
            "¿costolcl_previo1/2 respalda costo_local?",
            "En el Excel ambas filas mencionan costo_local; no hemos migrado costos por riesgo contable.",
        ],
    )

    question_block(
        doc, 16, "Historial cuando el mismo producto aparece 2 veces con precios distintos",
        [
            "La tabla costos es 1:1 con el producto.",
            "¿Les sirve un historial por línea de compra (Control 25)?",
        ],
    )

    question_block(
        doc, 17, "Decimales en costos",
        [
            "¿Costo con 4 decimales solo en compras y redondeo a 2 al pasar a costos?",
            "¿En qué momento se redondea: al confirmar compra, al promediar, al mostrar?",
        ],
    )

    question_block(doc, 18, "referencia_costo", ["¿Reducir de 50 a 30 caracteres o mantener 50?"])
    question_block(doc, 19, "Nodocto", ["¿Tipo y longitud del campo?"])
    question_block(doc, 20, "iva_pagado", ["¿Se mantiene el campo iva_pagado en la tabla costos?"])

    doc.add_page_break()

    # --- CONFIRMACIONES ---
    heading(doc, "D. Confirmaciones — decisiones ya implementadas (validar)", 1)
    para(
        doc,
        "Estas decisiones ya están en el sistema. Solo necesitamos que el cliente confirme "
        "que están bien o indique correcciones.",
        italic=True,
    )

    question_block(
        doc, 21, "Código de proveedor",
        [
            "¿Siguen usando códigos numéricos (275) o migran a tipados (PL0005 / PE0001)?",
            "Hoy el sistema acepta ambos formatos.",
        ],
    )

    question_block(
        doc, 22, "Direcciones de proveedor",
        [
            "¿2 bloques fijos en la ficha (como sus capturas)?",
            "¿O tabla direcciones compartida proveedores/clientes con tipo (facturación, entrega, bodega)?",
        ],
    )

    question_block(
        doc, 23, "Contactos de proveedor",
        ["¿2 contactos fijos (como captura) o tabla flexible para 2–3 contactos?"],
    )

    question_block(
        doc, 24, "Teléfonos y giros",
        ["¿Campos fijos en la ficha o tablas aparte reutilizables?"],
    )

    question_block(
        doc, 25, "Vínculo de cuentas bancarias",
        [
            "¿Vínculo polimórfico (tipo_entidad + id)?",
            "¿O dos FK (idproveedor, idcliente)?",
            "Hoy: tabla aparte solo para proveedores.",
        ],
    )

    question_block(
        doc, 26, "Catálogo contable — estructura",
        [
            "¿Una tabla con columna Gasto/Activo o dos tablas separadas?",
            "Hoy: una tabla + dos vistas (Gastos / Activos).",
        ],
    )

    question_block(
        doc, 27, "Catálogo contable — datos reales",
        ["¿Pueden entregar el catálogo definitivo del contador?", "Los 9 rubros actuales son ejemplos del documento."],
    )

    question_block(
        doc, 28, "Catálogo contable — condición IVA sugerida",
        ["¿Al elegir un rubro se debe precargar la condición de IVA?", "Ej.: Impuestos municipales → No sujeto."],
    )

    question_block(
        doc, 29, "Campos marcados con * en productos (Excel)",
        [
            "¿Entran en esta etapa o quedan para después?",
            "idcategoria de 6 a 10 caracteres",
            "Renombrados a Observa / Observa2 — ¿confirman esos nombres?",
        ],
    )

    doc.add_page_break()

    # --- Resumen ejecutivo ---
    heading(doc, "E. Resumen ejecutivo — mínimo para avanzar", 1)
    para(
        doc,
        "Si el tiempo es limitado, estas 4 respuestas desbloquean la siguiente fase de desarrollo:",
        bold=True,
    )

    table = doc.add_table(rows=5, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["#", "Pregunta", "Desbloquea"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True

    rows_data = [
        ("1", "Descuento: ¿informativo o se resta?", "Cálculo de líneas y totales"),
        ("3", "Estados y cuándo afecta inventario/costos/CxP", "Flujo del documento"),
        ("5–6", "Totales exactos + fiscal vs. contable", "Cierre de factura"),
        ("2", "Fechas obligatorias del encabezado", "Control fiscal mensual"),
    ]
    for ri, row in enumerate(rows_data, start=1):
        for ci, val in enumerate(row):
            table.rows[ri].cells[ci].text = val

    doc.add_paragraph()
    para(
        doc,
        "Las demás preguntas pueden responderse en una segunda ronda (importaciones, "
        "nota de crédito, aranceles, cuentas por pagar, bancos, costos).",
        italic=True,
    )

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("— Fin del documento —")
    fr.font.size = Pt(10)
    fr.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

    doc.save(OUT)
    print(f"Generado: {OUT}")


if __name__ == "__main__":
    main()
